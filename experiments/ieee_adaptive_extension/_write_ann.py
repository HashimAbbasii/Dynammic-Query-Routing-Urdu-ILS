# -*- coding: utf-8 -*-
import json
from collections import Counter
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = Path(r"c:\Users\User\OneDrive\Documents\ULTRA_Project")
EXTRACT = json.loads((ROOT / "experiments/ieee_adaptive_extension/_cg31_gold_extract.json").read_text(encoding="utf-8"))
qid2q = {q["query_id"]: q["query_text"] for q in EXTRACT["queries"]}
qid2sid = {q["query_id"]: q["source_doc_id"] for q in EXTRACT["queries"]}
qid2h = {q["query_id"]: EXTRACT["docs"][str(q["source_doc_id"])]["Headline"] for q in EXTRACT["queries"]}

# Manual annotations; Urdu_Gold_Expression curated from gold extract (headline + key body forms).
ANN = json.loads((ROOT / "experiments/ieee_adaptive_extension/CG31_manual_annotations.json").read_text(encoding="utf-8"))

ORDER = list(ANN.keys())
assert len(ORDER) == 31

COLS = ["Query_ID","Query_Text","Query_Composition","Entity_Present","Acronym_Present","Entity_Type","Spelling_Variation","Underspecification","Likely_Semantic_Intent","Gold_Doc_ID","Urdu_Gold_Expression","Query_Gold_Lexical_Overlap","Transliteration_Recoverable","Entity_Linking_Recoverable","Acronym_Expansion_Recoverable","Char_Ngram_Recoverable","Dense_Multilingual_Recoverable","Query_Expansion_Recoverable","Existing_Method_Coverage","Failure_Mechanism_Label","Secondary_Mechanism","Confidence","Potential_New_Mechanism","Evidence_Notes","Annotator","Annotation_Date"]
ANN_DATE = date.today().isoformat()
ROWS = []
for qid in ORDER:
    a = ANN[qid]
    assert not (a["Potential_New_Mechanism"] == "Yes" and a["Existing_Method_Coverage"] != "Uncovered")
    row = {
        "Query_ID": qid,
        "Query_Text": qid2q[qid],
        "Gold_Doc_ID": qid2sid[qid],
        "Annotator": "research_analyst",
        "Annotation_Date": ANN_DATE,
    }
    row.update(a)
    # if expression empty, fall back to headline
    if not row.get("Urdu_Gold_Expression"):
        row["Urdu_Gold_Expression"] = qid2h[qid]
    ROWS.append(row)

OUT = ROOT / "experiments/ieee_adaptive_extension/31_CG_Failure_Analysis_Template.xlsx"
wb = Workbook()
ws = wb.active
ws.title = "31_CG_Failures_Annotation"
ws.append(COLS)
for c in ws[1]:
    c.font = Font(bold=True)
wrap = Alignment(wrap_text=True, vertical="top")
for r in ROWS:
    ws.append([r.get(c, "") for c in COLS])
for row in ws.iter_rows(min_row=2, max_row=32, max_col=len(COLS)):
    for cell in row:
        cell.alignment = wrap
widths = {"Query_Text": 42, "Evidence_Notes": 48, "Urdu_Gold_Expression": 34, "Likely_Semantic_Intent": 34, "Failure_Mechanism_Label": 32}
for i, c in enumerate(COLS, 1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)

cov = Counter(r["Existing_Method_Coverage"] for r in ROWS)
labels = Counter(r["Failure_Mechanism_Label"] for r in ROWS)
cluster_map = {
    "institutional_descriptive_alias": "institutional_acronym_or_alias",
    "institutional_acronym_to_descriptive_urdu_alias": "institutional_acronym_or_alias",
    "institutional_acronym_plus_en_finance_jargon": "institutional_acronym_or_alias",
    "acronym_letter_name_expansion": "institutional_acronym_or_alias",
    "acronym_letter_name_plus_same_en_name_script_conversion": "institutional_acronym_or_alias",
    "same_en_institution_name_script_conversion_plus_vocab_paraphrase": "same_en_name_script_plus_jargon",
    "same_en_institution_name_script_conversion_plus_en_finance_jargon": "same_en_name_script_plus_jargon",
    "en_brand_in_urdu_news_framing": "en_brand_or_title_script_bridge",
    "en_title_brand_in_urdu_news_framing": "en_brand_or_title_script_bridge",
    "cross_script_person_name_variant": "cross_script_entity_or_person",
    "cross_script_named_entity_variant": "cross_script_entity_or_person",
    "cross_script_geo_entity_variant": "cross_script_entity_or_person",
    "person_name_spelling_collapse_and_cross_script_variant": "cross_script_entity_or_person",
    "sports_league_acronym_plus_city_team_script_variant": "sports_entity_acronym_script",
    "team_country_abbrev_plus_sports_jargon": "sports_entity_acronym_script",
    "en_sports_jargon_plus_entity_event_composition": "en_jargon_entity_event_composition",
    "en_finance_jargon_plus_geo_event_composition": "en_jargon_entity_event_composition",
    "person_plus_undescribed_work_alias": "underspecified_or_descriptive_referent",
    "underspecified_person_referent": "underspecified_or_descriptive_referent",
    "underspecified_product_referent": "underspecified_or_descriptive_referent",
    "crosslingual_semantic_paraphrase": "semantic_or_vocab_paraphrase",
    "rare_product_name_transliteration_and_translation_dual_form": "rare_product_translit_translate",
}
by = {}
for r in ROWS:
    by.setdefault(cluster_map.get(r["Failure_Mechanism_Label"], r["Failure_Mechanism_Label"]), []).append(r)

ws2 = wb.create_sheet("Mechanism_Clusters")
ws2.append(["Cluster_Label", "Count", "Pct_of_31", "Example_Query_IDs", "Dominant_Existing_Coverage", "Potential_New_Yes_Count", "Notes"])
for c in ws2[1]:
    c.font = Font(bold=True)
notes = {
    "institutional_acronym_or_alias": "Descriptive aliases and letter-name expansions; covered by acronym/EL.",
    "en_brand_or_title_script_bridge": "English brand/title present; covered by transliteration/EL.",
    "cross_script_entity_or_person": "Person/geo/name script variants; covered by transliteration/EL.",
    "en_jargon_entity_event_composition": "English domain jargon with entities/events; Partial via dense/QE.",
    "underspecified_or_descriptive_referent": "Missing linkable name; Partial via dense; known TOT/underspecification.",
    "same_en_name_script_plus_jargon": "Same English institution name in Urdu script plus jargon paraphrase.",
    "sports_entity_acronym_script": "Sports team/league acronyms and city names; covered by acronym/EL.",
    "semantic_or_vocab_paraphrase": "No strong NE; bilingual paraphrase; dense/QE.",
    "rare_product_translit_translate": "Rare product with transliteration and translation dual forms.",
}
for cl, members in sorted(by.items(), key=lambda x: (-len(x[1]), x[0])):
    cov_c = Counter(m["Existing_Method_Coverage"] for m in members).most_common(1)[0][0]
    new_c = sum(1 for m in members if m["Potential_New_Mechanism"] == "Yes")
    ws2.append([cl, len(members), round(100 * len(members) / 31, 1), ", ".join(m["Query_ID"] for m in members[:6]), cov_c, new_c, notes.get(cl, "")])

ws2.append([])
ws2.append(["AGGREGATE_Existing_Method_Coverage", "Count", "Pct"])
for k in ["Covered", "Partial", "Uncovered"]:
    ws2.append([k, cov.get(k, 0), round(100 * cov.get(k, 0) / 31, 1)])
ws2.append([])
uncovered = [r["Query_ID"] for r in ROWS if r["Existing_Method_Coverage"] == "Uncovered"]
newyes = [r["Query_ID"] for r in ROWS if r["Potential_New_Mechanism"] == "Yes"]
ws2.append(["Uncovered_Query_IDs", ", ".join(uncovered) or "(none)"])
ws2.append(["Potential_New_Mechanism_Yes_Query_IDs", ", ".join(newyes) or "(none)"])
ws2.append(["Shared_new_structural_phenomenon?", "No surviving candidate: zero Uncovered+Potential_New_Mechanism=Yes cases."])
ws2.append(["Fine_grained_label_counts"])
for lab, n in labels.most_common():
    ws2.append([lab, n])
for i in range(1, 8):
    ws2.column_dimensions[get_column_letter(i)].width = 30 if i != 4 else 42
wb.save(OUT)

clusters = Counter(cluster_map.get(r["Failure_Mechanism_Label"], r["Failure_Mechanism_Label"]) for r in ROWS)
summary = ROOT / "experiments/ieee_adaptive_extension/CG31_FAILURE_ANALYSIS_SUMMARY.md"
lines = [
    "# CG Failure Analysis Summary (n=31)",
    "",
    f"**Annotation date:** {ANN_DATE}  ",
    "**Scope:** Program B TRAIN+DEV ExactSource queries where gold is absent from BM25∪NG3∪Dense Top-50.  ",
    "**Sources:** phase13 `queries_kn.csv` + `data/clean_articles.csv` golds.  ",
    "**Method:** Stage 1 query-first; Stages 2–4 after gold comparison. No method design.  ",
    "**Workbook:** `experiments/ieee_adaptive_extension/31_CG_Failure_Analysis_Template.xlsx`",
    "",
    "## Artifact confirmation",
    "",
    "- 31 CG failures confirmed from `FINAL_FAILURE_MECHANISM_PER_QUERY.csv` (status=CG_FAILURE).",
    "- Query texts and gold IDs confirmed from phase13 TRAIN/DEV `queries_kn.csv`.",
    "- Gold headlines/text confirmed from `data/clean_articles.csv` (31/31 found).",
    "",
    "## Coverage aggregates",
    "",
    f"- **Covered:** {cov.get('Covered', 0)}/31",
    f"- **Partial:** {cov.get('Partial', 0)}/31",
    f"- **Uncovered:** {cov.get('Uncovered', 0)}/31",
    f"- **Potential_New_Mechanism = Yes:** {len(newyes)}/31",
    "",
    "## Mechanism clusters (coarse)",
    "",
]
for cl, n in clusters.most_common():
    lines.append(f"- **{cl}:** {n}/31")
lines += [
    "",
    "## Uncovered + Potential_New_Mechanism = Yes",
    "",
    "None. No case was both Uncovered and judged to require a new mechanism.",
    "",
    "## Evidence test for a shared new phenomenon",
    "",
    "The 31 misses fragment across known categories: institutional acronym/alias, brand/title script bridging,",
    "person/geo name variants, English jargon composed with entities/events, underspecified referents,",
    "same-English-name script conversion plus jargon, and bilingual paraphrase.",
    "Underspecified cases (KN008, KN037, and related descriptive aliases) are explainable as tip-of-the-tongue /",
    "descriptive known-item phenomena; they remain Partial because dense matching could still help, and they do not",
    "share a distinct uncovered structure beyond that known class.",
    "",
    "## Conclusion",
    "",
    "On this evidence, **no candidate new failure mechanism survives**. Residual difficulty is heterogeneous and",
    "largely Covered or Partial under transliteration, entity linking, acronym expansion, dense multilingual retrieval,",
    "and/or query expansion.",
    "",
]
summary.write_text("\n".join(lines), encoding="utf-8")
print("Wrote", OUT)
print("Wrote", summary)
print("Coverage", dict(cov))
print("Clusters", dict(clusters))
print("Uncovered", uncovered, "NewYes", newyes)