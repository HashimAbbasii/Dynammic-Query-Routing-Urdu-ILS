#!/usr/bin/env python3
"""Build ULTRA advisor presentation (read-only of frozen numbers)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ULTRA_Advisor_Presentation.pptx"
FIG = ROOT / "Papers" / "PLOS_ONE" / "figures"

NAVY = RGBColor(0x0B, 0x3D, 0x5C)
NAVY_DARK = RGBColor(0x07, 0x28, 0x3C)
GOLD = RGBColor(0xC4, 0x96, 0x2C)
TEAL = RGBColor(0x1F, 0x7A, 0x6B)
CRIMSON = RGBColor(0xA6, 0x3D, 0x40)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF = RGBColor(0xF7, 0xF4, 0xEE)
SLATE = RGBColor(0x2F, 0x3A, 0x43)
MUTED = RGBColor(0x5C, 0x68, 0x72)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD7, 0xDE, 0xE4)
GREEN = RGBColor(0x2E, 0x6B, 0x4F)

W = Inches(13.333)
H = Inches(7.5)


def rgb_hex(c: RGBColor) -> str:
    return f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"


def set_run(run, text, size=18, bold=False, color=SLATE, name="Calibri", italic=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = name


def add_text(slide, l, t, w, h, text, size=18, bold=False, color=SLATE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, name="Calibri"):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    anchor_map = {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}
    tf._txBody.bodyPr.set("anchor", anchor_map[anchor])
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    set_run(r, text, size, bold, color, name, italic)
    return box


def fill_tf(tf, lines, size=17, color=SLATE, bold=False, spacing=10, bullet=True):
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        p.level = 0
        if isinstance(line, tuple):
            text, is_bold = line
        else:
            text, is_bold = line, bold
        prefix = "•  " if bullet else ""
        r = p.add_run()
        set_run(r, prefix + text, size, is_bold, color)
    return tf


def shape_fill(shape, color, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)


def rect(slide, l, t, w, h, color, line=None, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(l), Inches(t), Inches(w), Inches(h))
    shape_fill(s, color, line)
    if rounded:
        try:
            s.adjustments[0] = 0.08
        except Exception:
            pass
    return s


def notes(slide, text):
    ns = slide.notes_slide.notes_text_frame
    ns.text = text


def footer(slide, page, total, backup=False):
    rect(slide, 0, 7.28, 13.333, 0.22, NAVY_DARK)
    label = "BACKUP  ·  " if backup else ""
    add_text(slide, 0.45, 7.28, 9.5, 0.22, f"{label}ULTRA  ·  Advisor review  ·  Hashim Shazad  ·  14 Sep 2026", 10, False, RGBColor(0xC5, 0xD0, 0xD8), PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    add_text(slide, 11.2, 7.28, 1.7, 0.22, f"{page}  /  {total}", 10, False, GOLD, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)


def header(slide, kicker, title, subtitle=None):
    rect(slide, 0, 0, 13.333, 1.18, NAVY)
    rect(slide, 0, 0, 0.12, 7.5, GOLD)
    add_text(slide, 0.45, 0.12, 12.4, 0.28, kicker.upper(), 11, True, GOLD, name="Calibri")
    add_text(slide, 0.45, 0.36, 12.4, 0.46, title, 26, True, WHITE, name="Calibri")
    if subtitle:
        add_text(slide, 0.45, 0.82, 12.4, 0.28, subtitle, 13, False, RGBColor(0xC9, 0xD6, 0xDE))


def card(slide, l, t, w, h, title, body_lines, accent=TEAL, title_size=14, body_size=14):
    rect(slide, l, t, w, h, WHITE, LINE, rounded=True)
    rect(slide, l, t, 0.10, h, accent)
    add_text(slide, l + 0.28, t + 0.12, w - 0.4, 0.36, title, title_size, True, NAVY)
    box = slide.shapes.add_textbox(Inches(l + 0.28), Inches(t + 0.48), Inches(w - 0.45), Inches(h - 0.62))
    fill_tf(box.text_frame, body_lines, size=body_size, color=SLATE, spacing=6)
    return box


def kpi(slide, l, t, w, h, value, label, sub, accent=GOLD):
    rect(slide, l, t, w, h, WHITE, LINE, rounded=True)
    rect(slide, l, t, w, 0.08, accent)
    add_text(slide, l + 0.15, t + 0.22, w - 0.3, 0.7, value, 32, True, NAVY, PP_ALIGN.CENTER)
    add_text(slide, l + 0.15, t + 0.92, w - 0.3, 0.4, label, 14, True, accent, PP_ALIGN.CENTER)
    add_text(slide, l + 0.18, t + 1.32, w - 0.36, 0.7, sub, 12, False, MUTED, PP_ALIGN.CENTER)


def style_table(table, header=True, font=12):
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    run.font.size = Pt(font)
                    run.font.name = "Calibri"
                    run.font.bold = i == 0 if header else False
                    run.font.color.rgb = WHITE if i == 0 and header else SLATE
            fill = NAVY if i == 0 and header else (RGBColor(0xEE, 0xF3, 0xF6) if i % 2 else WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill


def set_cell(table, r, c, text, bold=False):
    table.cell(r, c).text = text
    for p in table.cell(r, c).text_frame.paragraphs:
        for run in p.runs:
            run.font.bold = bold or r == 0


def blank(prs, page, total, backup=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 13.333, 7.5, OFF)
    footer(slide, page, total, backup)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    total = 20

    # 1 Title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, NAVY_DARK)
    rect(s, 0, 0, 0.18, 7.5, GOLD)
    add_text(s, 0.7, 1.15, 12, 0.35, "AIR UNIVERSITY  ·  MS THESIS  /  ADVISOR REVIEW", 13, True, GOLD)
    add_text(s, 0.7, 1.6, 12, 1.3, "ULTRA", 60, True, WHITE)
    add_text(s, 0.7, 2.85, 12, 1.0, "Script-aware retrieval for Urdu and Roman Urdu news", 26, False, RGBColor(0xD7, 0xE2, 0xEA))
    add_text(s, 0.7, 3.9, 11.5, 0.7, "What was built, what the frozen numbers actually mean,\nand what should happen next — without inflating a headline.", 16, False, RGBColor(0xB7, 0xC4, 0xCD))
    rect(s, 0.7, 4.85, 2.2, 0.06, GOLD)
    add_text(s, 0.7, 5.15, 11, 0.35, "Hashim Shazad   ·   Advisor: Adnan Aslam   ·   14 September 2026", 15, False, WHITE)
    add_text(s, 0.7, 5.55, 11, 0.35, "Frozen system M0  ·  PLOS ONE / IEEE packaging  ·  ULTRA v2 TRAIN/DEV (TEST sealed)", 13, False, RGBColor(0x9A, 0xAB, 0xB6))
    notes(s, "Open by saying: this is not a 87% accuracy talk. It is a freeze, three different evaluations, and a Roman-Urdu follow-on program. Ask the advisor to hold questions until the decision slide if time is tight (20 minutes).")

    # 2 30-second
    s = blank(prs, 2, total)
    header(s, "If you remember one slide", "The 30-second version", "Three different questions. One frozen system. One remaining problem.")
    kpi(s, 0.45, 1.5, 4.0, 2.25, "87.18%", "Development known-item", "68/78 ExactSource Hit@5\nTitle-like QTRN pool — not unseen chat", GOLD)
    kpi(s, 4.65, 1.5, 4.0, 2.25, "67.50%", "New sealed known-item", "27/40 on K001–K040\nUrdu 26/28 · Roman 1/12", TEAL)
    kpi(s, 8.85, 1.5, 4.0, 2.25, "57.50%", "New human usefulness", "23/40 Success@5 on U\nUrdu 17/18 · Roman 6/18", CRIMSON)
    card(s, 0.45, 4.0, 6.2, 2.95, "What this project already is", [
        "A frozen, reproducible lexical system (M0) — not the old SVM/MiniLM router.",
        "A PLOS/IEEE manuscript that reports three numbers and does not average them.",
        "A documented finding: Urdu-script news search is strong; ordinary Roman Urdu is not.",
    ], TEAL, 15, 15)
    card(s, 6.9, 4.0, 5.95, 2.95, "What it is not", [
        "Not a published DOI in the repo (submission snapshot).",
        "Not 80% real-world accuracy.",
        "Not finished Roman retrieval — that is ULTRA v2, TEST still sealed.",
    ], CRIMSON, 15, 15)
    notes(s, "Say the three numbers out loud and the three questions. Then: we will not average them. Roman is the gap. v2 exists to attack that gap without rewriting the paper.")

    # 3 Agenda
    s = blank(prs, 3, total)
    header(s, "Meeting plan", "What I need from this review", "I am not asking to reopen frozen M0. I am asking for a decision on thesis vs follow-on.")
    items = [
        ("1", "What I did", "Froze M0, sealed K/U, human labels, ablation, papers, thesis draft."),
        ("2", "What the numbers mean", "87% is development known-item. 67.5% is new known-item. 57.5% is usefulness."),
        ("3", "What v2 already showed", "On 51 new Roman KN queries, Method D collapses; dense helps; fusion/letter-name failed."),
        ("4", "The ask", "Finish thesis + PLOS on M0. Keep v2 TEST sealed. Treat v2 as a second paper, not a rewrite."),
    ]
    y = 1.5
    for n, t, b in items:
        rect(s, 0.5, y, 12.3, 1.25, WHITE, LINE, True)
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.75), Inches(y + 0.32), Inches(0.6), Inches(0.6))
        shape_fill(circ, NAVY)
        add_text(s, 0.75, y + 0.38, 0.6, 0.48, n, 16, True, WHITE, PP_ALIGN.CENTER)
        add_text(s, 1.6, y + 0.18, 10.8, 0.4, t, 18, True, NAVY)
        add_text(s, 1.6, y + 0.62, 10.8, 0.45, b, 15, False, MUTED)
        y += 1.35
    notes(s, "This is the contract for the meeting. If time is short, jump from this slide to the three numbers, then Roman split, then the ask.")

    # 4 Problem
    s = blank(prs, 4, total)
    header(s, "The research problem", "Urdu users do not type in one script", "Corpus: 111,860 news articles. A single native-script index misses Roman queries even when the article exists.")
    card(s, 0.45, 1.5, 4.05, 5.4, "Perso-Arabic Urdu", [
        "Native script of the collection.",
        "BM25 on article text is a strong known-item retriever for title-like Urdu queries.",
        "On sealed K: 26/28 Hit@5.",
        "On sealed U: 17/18 Success@5.",
    ], TEAL, 16, 15)
    card(s, 4.65, 1.5, 4.05, 5.4, "Roman Urdu", [
        "Latin letters, informal spelling, no standard orthography.",
        "Development Roman was title_roman (our own romanizer) — Method D looked excellent (22/23).",
        "Ordinary Roman titles on K: 1/12 Hit@5.",
        "Chat-style Roman on U: 6/18 Success@5.",
    ], CRIMSON, 16, 15)
    card(s, 8.85, 1.5, 4.0, 5.4, "Mixed / bilingual", [
        "Users mix scripts and English entities.",
        "M0 sends MIXED to Urdu BM25.",
        "U mixed cell: 0/4 (descriptive, small n).",
        "Not a solved path. Do not hide it; do not over-generalize it.",
    ], GOLD, 16, 15)
    notes(s, "Frame the problem as bilingual IR, not as 'we built a chatbot'. The scientific contribution is a frozen protocol that keeps these scripts separate.")

    # 5 SVM rejected
    s = blank(prs, 5, total)
    header(s, "Honesty first", "We rejected the tempting official story", "Original thesis Layer A: SVM routes queries to MiniLM headline vs full-article indexes.")
    add_text(s, 0.5, 1.45, 12.3, 0.4, "On held-out trap queries H001–H040 the router could match labels and still hurt retrieval.", 16, False, SLATE)
    tbl = s.shapes.add_table(4, 3, Inches(0.5), Inches(2.05), Inches(12.3), Inches(2.4)).table
    tbl.columns[0].width = Inches(5.2)
    tbl.columns[1].width = Inches(3.5)
    tbl.columns[2].width = Inches(3.6)
    data = [
        ["Layer", "SVM", "Word-count / other"],
        ["Classification (n=50, Phase 3B)", "86%", "84%"],
        ["Held-out trap classification (n=40)", "60%", "20%"],
        ["Held-out dual-index P@5 (400 judgments)", "33.00%", "36.50%  (SVM lost)"],
    ]
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            set_cell(tbl, r, c, val, r == 0)
    style_table(tbl, True, 14)
    card(s, 0.5, 4.7, 12.3, 2.2, "Lesson we kept for the rest of the project", [
        "Classification accuracy is not an IR result. Older drafts quoting ~90% P@15 or 100% routing are Layer A — not M0.",
        "Official system became script-aware BM25. MiniLM dual-index work is archived, not the paper headline.",
        "That is a scientific strength: we stopped when retrieval disagreed with the protocol labels.",
    ], GOLD, 15, 15)
    notes(s, "If the advisor still thinks of the SVM thesis, this is the correction slide. Do not apologize — show we chose the retrieval metric over a prettier classification number.")

    # 6 M0 architecture
    s = blank(prs, 6, total)
    header(s, "Official frozen system — M0", "Unicode detector + two BM25 indexes", "Frozen 27 Aug 2026. SVM = false. k1=1.5, b=0.75. Top-50 retrieve, Top-5 official cutoff.")
    img = FIG / "Fig1_m0_routing.png"
    s.shapes.add_picture(str(img), Inches(0.4), Inches(1.45), Inches(8.3), Inches(3.55))
    card(s, 8.85, 1.45, 4.0, 5.45, "What is frozen", [
        "URDU / MIXED / OTHER → Urdu BM25 on article text.",
        "ROMAN → Method D: romanize documents, search the query as typed.",
        "198-key dictionary + character table. Dictionary keys were not edited after freeze.",
        "Corpus SHA-256 verified locally (111,860 docs).",
        "Frozen code: run_phase5.py (phase5_roman_urdu).",
        "This is the PLOS/IEEE system. v2 must not rewrite it.",
    ], TEAL, 14, 13)
    notes(s, "Walk the figure left to right. Emphasize MIXED goes to Urdu BM25 by design. Method D is document-side romanization, not a neural model.")

    # 7 Development
    s = blank(prs, 7, total)
    header(s, "Development / validation  ·  n = 78", "Why M0 was selected: 68/78 = 87.18%", "Same pool as the bar chart. Roman subset was title_roman — Method D’s own spelling family.")
    s.shapes.add_picture(str(FIG / "Fig2_development_comparators.png"), Inches(0.35), Inches(1.4), Inches(8.4), Inches(5.5))
    card(s, 8.85, 1.45, 4.0, 5.45, "How to say this number", [
        "True development known-item Hit@5.",
        "Urdu-only BM25 on the same pool: 0.5897 — the Roman path was necessary here.",
        "Method D recovered 22/23 development Roman items.",
        "Forbidden sentence: “the system achieves 87% on unseen queries.”",
        "Allowed sentence: “On the freeze known-item pool, ExactSource Hit@5 = 68/78.”",
    ], GOLD, 14, 13)
    notes(s, "Point at MiniLM bars: we already tried dense/chunking on this pool; they lost to M0 here because development Roman matched Method D. That advantage did not transfer to ordinary Roman.")

    # 8 Three numbers
    s = blank(prs, 8, total)
    header(s, "Official frozen table", "Three numbers. Three questions. Never one average.", "Copied from sealed Phase 8–12 reports. Not recomputed for this slide.")
    rows = [
        ("87.18%", "68/78", "Development known-item", "Did the designated source reach Top-5 on title-like QTRN queries?", GOLD),
        ("67.50%", "27/40", "Sealed known-item K", "Same question, new headlines. Ordinary Roman titles, not title_roman.", TEAL),
        ("57.50%", "23/40", "Sealed human U", "Was anything in Top-5 useful (A or B)? No gold article.", CRIMSON),
    ]
    y = 1.45
    for val, frac, name, q, acc in rows:
        rect(s, 0.45, y, 12.4, 1.55, WHITE, LINE, True)
        rect(s, 0.45, y, 0.12, 1.55, acc)
        add_text(s, 0.8, y + 0.18, 2.6, 0.7, val, 32, True, NAVY)
        add_text(s, 0.8, y + 0.88, 2.6, 0.4, frac, 14, True, acc)
        add_text(s, 3.6, y + 0.22, 8.9, 0.45, name, 20, True, NAVY)
        add_text(s, 3.6, y + 0.75, 8.9, 0.55, q, 15, False, MUTED)
        y += 1.7
    add_text(s, 0.5, 6.7, 12.3, 0.4, "Diagnostic only — do not quote as official unseen: H001–H040 Success@5 = 25/40 = 62.5% (traps; burned).", 13, False, MUTED)
    notes(s, "This is the thesis Table 1. Pause. If the advisor asks for one number, refuse politely and repeat the three questions.")

    # 9 Roman failure
    s = blank(prs, 9, total)
    header(s, "The scientific finding", "Urdu-script search works. Ordinary Roman Urdu does not.", "Drop 87.18% → 67.50% is concentrated on Roman titles, not on Urdu BM25 breaking.")
    s.shapes.add_picture(str(FIG / "Fig3_script_splits.png"), Inches(0.3), Inches(1.4), Inches(8.6), Inches(5.5))
    card(s, 8.95, 1.45, 3.9, 5.45, "What to claim", [
        "K Urdu titles: 26/28 Hit@5.",
        "K Roman titles: 1/12 Hit@5.",
        "U Urdu needs: 17/18 Success@5.",
        "U Roman needs: 6/18 Success@5.",
        "U mixed: 0/4 (small n).",
        "Honest one-liner for the thesis: native-script news search is strong; ordinary Roman Urdu remains the main failure mode.",
    ], CRIMSON, 14, 13)
    notes(s, "This is the slide that convinces a careful advisor you understand your own result. Development 22/23 Roman was the wrong Roman.")

    # 10 Candidate generation
    s = blank(prs, 10, total)
    header(s, "Why Roman fails", "It is mostly a candidate-generation problem, not a Top-5 rerank problem", "On sealed K, most Roman sources never enter the Top-50. Reranking cannot recover a document that was never retrieved.")
    s.shapes.add_picture(str(FIG / "Fig4_k_miss_analysis.png"), Inches(0.3), Inches(1.4), Inches(8.5), Inches(5.5))
    card(s, 8.9, 1.45, 3.95, 5.45, "Implication for next work", [
        "Urdu misses are rare and often still in ranks 6–50.",
        "Roman: 10 of 12 sources absent from Top-50.",
        "That is why Phase 11 query expansions did not beat 68/78, and why v2 letter-name / 3-way RRF failed later.",
        "Next method must get Roman golds into the pool — dense, better romanization, or a new matching feature.",
    ], GOLD, 14, 13)
    notes(s, "This slide justifies ULTRA v2. Do not propose another BM25 stoplist. The gold is not in the list.")

    # 11 Integrity
    s = blank(prs, 11, total)
    header(s, "Why this is defendable", "The evaluation protocol is the contribution, not a rounded percentage", "We can show a reviewer exactly what was frozen, what was sealed, and what was burned.")
    bits = [
        ("Freeze", "27 Aug 2026 manifest. Do not change k1, b, detector, Method D, corpus, dictionary keys."),
        ("Sealed tests", "K and U sealed before retrieval. Not QTRN, not H. Official unseen tests of this M0."),
        ("Ablation", "M1–M4 query-side Roman expansions: all still 68/78. M0 not replaced."),
        ("Human labels", "U A1 official 23/40. Independent A2: 26/40; five-way κ=0.549; binary κ=0.682. A2 does not replace A1."),
        ("Forbidden claim", "~80% usefulness is written in the protocol as rejected. We did not chase 90% Hit@5 after freeze."),
        ("Burned sets", "K, U, and H cannot be used to retune M0, Method D, or the dictionary."),
    ]
    positions = [(0.45, 1.45), (4.55, 1.45), (8.65, 1.45), (0.45, 4.2), (4.55, 4.2), (8.65, 4.2)]
    colors = [TEAL, TEAL, GOLD, GOLD, CRIMSON, CRIMSON]
    for (l, t), (title, body), col in zip(positions, bits, colors):
        rect(s, l, t, 3.95, 2.5, WHITE, LINE, True)
        rect(s, l, t, 3.95, 0.08, col)
        add_text(s, l + 0.2, t + 0.22, 3.55, 0.4, title, 16, True, NAVY)
        add_text(s, l + 0.2, t + 0.7, 3.55, 1.55, body, 13, False, SLATE)
    notes(s, "This is how you convince a methods-strict advisor. We have hashes, seals, and a negative ablation. That is more valuable than a 90% claim.")

    # 12 Deliverables
    s = blank(prs, 12, total)
    header(s, "Written output", "Papers and thesis — status in the repository", "Repo evidence only: no DOI or acceptance letter was found.")
    tbl = s.shapes.add_table(6, 3, Inches(0.4), Inches(1.45), Inches(12.5), Inches(4.55)).table
    tbl.columns[0].width = Inches(3.3)
    tbl.columns[1].width = Inches(4.6)
    tbl.columns[2].width = Inches(4.6)
    data = [
        ["Manuscript", "What it is", "Status in repo"],
        ["PLOS ONE", "Script-aware BM25 for Urdu and Roman Urdu news search", "Editorial Manager package 6 Sep 2026. Snapshot branch publication/plos-one-final. Not proven published."],
        ["IEEE (M0)", "Same three official metrics, IEEE style", "Submission ZIP exists. Separate venue, not a new system."],
        ["AU MS thesis", "Word + markdown draft aligned to frozen M0", "Live file under Thesis_template/FINAL/. Layer A still labeled historical."],
        ["Historical MiniLM IEEE", "SVM / dual-index draft", "Archived. Do not present as the official result."],
        ["ULTRA v2 paper", "Roman first-stage follow-on", "No manuscript yet. Experiment reports only. TEST sealed."],
    ]
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            set_cell(tbl, r, c, val, r == 0)
    style_table(tbl, True, 12)
    for r in range(1, 6):
        for c in range(3):
            for p in tbl.cell(r, c).text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
    add_text(s, 0.5, 6.2, 12.3, 0.7, "Authors on the M0 papers: Hashim Shazad, Adnan Aslam. Independent U labels (A2): Areena Rahman.", 13, False, MUTED)
    notes(s, "If asked 'is it published?': the repo shows a submission-ready PLOS package dated 6 Sep 2026. Confirm portal status verbally if you have it; do not invent a DOI.")

    # 13 Why v2
    s = blank(prs, 13, total)
    header(s, "ULTRA v2  ·  research/ultra-v2-strengthening", "A follow-on program — not a rewrite of PLOS", "New TRAIN / DEV / TEST benchmark. K/U/H are burned for v2. TEST query files have not been opened.")
    card(s, 0.45, 1.45, 6.2, 5.45, "Why a new benchmark was required", [
        "K and U were sealed, retrieved, labeled, sliced by script, and put in the PLOS trail.",
        "Using them to invent new Roman rules would be test leakage.",
        "Phase 12 already said: if the system changes, K/U are burned. v2 is a system-change program.",
        "Objective: any improvement over frozen M0 must survive a genuinely unseen TEST.",
        "~80% is an aspiration, not a pass/fail gate.",
    ], TEAL, 15, 14)
    card(s, 6.9, 1.45, 5.95, 5.45, "What has been scored so far", [
        "Roman known-item TRAIN+DEV only: n = 51.",
        "NL track not scored. TEST not retrieved.",
        "Same corpus and dictionary SHAs as M0.",
        "Method D is imported read-only from run_phase5.py.",
        "PLOS numbers stay historical. They are never averaged into a v2 headline.",
    ], GOLD, 15, 14)
    notes(s, "Make the relationship explicit: v2 is meant to improve Roman retrieval for a later paper. It does not replace 87/67.5/57.5.")

    # 14 v2 scoreboard
    s = blank(prs, 14, total)
    header(s, "v2 TRAIN+DEV  ·  Roman KN n = 51  ·  ExactSource", "Scoreboard — dense is the only large Hit@5 move so far", "TEST sealed. These numbers are not a new official system.")
    tbl = s.shapes.add_table(7, 7, Inches(0.35), Inches(1.4), Inches(12.6), Inches(4.15)).table
    widths = [3.3, 1.5, 1.3, 1.3, 1.3, 1.4, 2.5]
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    data = [
        ["Method", "Decision", "Hit@1", "Hit@5", "Hit@10", "Hit@50", "MRR"],
        ["Method-D BM25 (R2-B0)", "Baseline", "1", "4", "4", "6", "0.0375"],
        ["Dense e5-small", "Partial", "5", "15", "16", "22", "0.1726"],
        ["Hybrid RRF k=60", "Partial", "3", "11", "19", "25", "0.1422"],
        ["R2-NG3 char 3-grams", "Partial", "3", "6", "6", "8", "0.0750"],
        ["Letter-name (Phase 7)", "Unsupported", "0", "0", "0", "1", "0.0007"],
        ["3-way RRF k=60", "Unsupported", "4", "8", "14", "24", "0.1310"],
    ]
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            set_cell(tbl, r, c, val, r == 0)
    style_table(tbl, True, 12)
    # highlight dense hit@5 cell
    cell = tbl.cell(2, 3)
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(0xD9, 0xEE, 0xE6)
    add_text(s, 0.45, 5.7, 12.4, 1.2, "Reading: Method D, which produced 22/23 on development title_roman, is 4/51 on new Roman KN. Dense is the strongest first-stage (15/51 Hit@5). Hybrid helps Hit@10/@50 but loses to dense at Hit@5. Letter-name and 3-way RRF are closed.", 14, False, SLATE)
    notes(s, "Say 4/51 vs 22/23 slowly. That is the generalization failure of Method D. Then: we already tried the obvious patches; two of them are unsupported.")

    # 15 Closed doors
    s = blank(prs, 15, total)
    header(s, "Controlled negatives are progress", "What we should stop repeating", "Each of these was preregistered, run once, and written up. That is how we avoid wasting the next month.")
    cards = [
        (CRIMSON, "Closed — do not rerun", [
            "Document-side Hunterian fallback (R2-1): no Top-50 gain, ROOM 0/11.",
            "Query-side letter-name expansion (Phase 7): 0/23 remaining misses; destroyed 5 of 6 BM25 Top-50 hits.",
            "Unweighted 3-way RRF of BM25+dense+NG3 (Phase 8): Hit@5 11→8 vs 2-way hybrid. Dilutes NG3 ranks.",
        ]),
        (GOLD, "Partial — keep as evidence", [
            "Dense e5-small: recovered 20 BM25 Top-50 misses; lost all 4 BM25 Hit@5 successes. Complementary, not a drop-in replacement.",
            "2-way hybrid: restores BM25 recalls; still 25 dual-misses that neither list contains.",
            "NG3: recovered KN035 (rank 5) and KN050 (rank 1) — two dual-miss golds. Not a general Roman solver.",
        ]),
        (TEAL, "Still true constraints", [
            "23 queries remain misses for BM25 ∩ dense ∩ NG3 ∩ hybrid.",
            "Routing on this set: 51/51 ROMAN — the detector is not the bug.",
            "Opening TEST now would measure a still-weak Roman KN system. Fine for a failure paper; wrong if we want a 'we beat M0' claim.",
        ]),
    ]
    x = 0.4
    for col, title, lines in cards:
        card(s, x, 1.45, 4.15, 5.45, title, lines, col, 15, 13)
        x += 4.3
    notes(s, "This is the 'I will not waste your time' slide. You already know letter-name and 3-way fusion failed. Ask for permission to stop those lines.")

    # 16 Recommend
    s = blank(prs, 16, total)
    header(s, "Recommended plan", "What I should do next", "Recommended option is B: finish the degree on frozen M0; keep v2 as a follow-on paper.")
    recs = [
        ("A  ·  Now (this month)", GREEN, [
            "Keep M0 frozen. Do not retune on K, U, or H.",
            "Thesis Table 1 = 87.18% / 67.50% / 57.50% with the three questions written above the numbers.",
            "Submit or track PLOS ONE; IEEE is the same science in another template.",
            "Strip any remaining ~90% P@15 / 100% routing language from the live thesis if it still appears unlabeled.",
        ]),
        ("B  ·  ULTRA v2 (follow-on paper)", TEAL, [
            "Do not open TEST until one system is preregistered and frozen on TRAIN/DEV.",
            "Do not delay the MS thesis waiting for v2 TEST.",
            "If one more TRAIN/DEV experiment: attack candidate generation (dense / better Roman matching) — not another unweighted RRF, not letter-name.",
            "Frame the next paper as Roman Urdu first-stage retrieval, not as beating 87%.",
        ]),
        ("C  ·  Only if you want a delay", CRIMSON, [
            "Hold the thesis until v2 TEST. I do not recommend this.",
            "Reason: TEST is sealed, NL is unlabeled, current best Roman KN Hit@5 is 15/51.",
            "That is a research program, not a graduation blocker. M0 is already a complete, honest study.",
        ]),
    ]
    x = 0.4
    for title, col, lines in recs:
        card(s, x, 1.45, 4.15, 5.45, title, lines, col, 14, 13)
        x += 4.3
    notes(s, "Recommend B explicitly. If the advisor loves v2, still freeze the thesis on M0. v2 can be a second paper after the degree, or a chapter labeled exploratory.")

    # 17 What not to do
    s = blank(prs, 17, total)
    header(s, "Risk control", "What we should not do", "These would weaken the thesis, the PLOS paper, or the v2 TEST.")
    bads = [
        ("Do not average 87% with 57%", "They answer different questions. A reviewer will treat that as overclaiming."),
        ("Do not open v2 TEST to “see how we are doing”", "One look burns the only unseen split. Seal metadata only."),
        ("Do not tune on K / U / H", "Already retrieved, labeled, and published. Leakage."),
        ("Do not claim ~80% usefulness", "Written as a rejected sentence in our own protocol."),
        ("Do not rerun letter-name or 3-way RRF", "Unsupported under preregistration. Need a new hypothesis."),
        ("Do not overwrite PLOS numbers with v2 TRAIN/DEV", "4/51 or 15/51 is not a replacement for 68/78."),
    ]
    y = 1.45
    for i, (t, b) in enumerate(bads):
        col = 0.45 if i % 2 == 0 else 6.9
        if i % 2 == 0 and i > 0:
            y += 1.75
        rect(s, col, y, 6.2, 1.6, WHITE, LINE, True)
        rect(s, col, y, 0.1, 1.6, CRIMSON)
        add_text(s, col + 0.28, y + 0.15, 5.7, 0.45, t, 15, True, NAVY)
        add_text(s, col + 0.28, y + 0.65, 5.7, 0.75, b, 13, False, MUTED)
    notes(s, "Read two of these aloud: averaging, and opening TEST. Those are the two ways this project gets scientifically damaged.")

    # 18 Decision
    s = blank(prs, 18, total)
    header(s, "The ask", "Decision I need from the advisor", "My recommendation is highlighted.")
    add_text(s, 0.5, 1.4, 12.3, 0.4, "Please pick one. I will follow it without mixing PLOS freeze and v2 TEST.", 15, False, SLATE)
    opts = [
        (False, "1", "Thesis + PLOS only", "Park v2. Graduate on frozen M0. Roman failure is the honest limitation section."),
        (True, "2", "Recommended: M0 thesis + v2 as follow-on", "Freeze the degree on the three numbers. Continue v2 on a new paper. TEST stays sealed until one candidate is frozen."),
        (False, "3", "Delay thesis for v2 TEST", "Not recommended. TEST unlabeled NL, Roman KN still weak on TRAIN/DEV, risk of rushing a leak."),
    ]
    y = 1.95
    for rec, n, title, body in opts:
        bg = RGBColor(0xE7, 0xF4, 0xEE) if rec else WHITE
        acc = GREEN if rec else (CRIMSON if n == "3" else NAVY)
        rect(s, 0.5, y, 12.3, 1.45, bg, acc if rec else LINE, True)
        add_text(s, 0.75, y + 0.18, 1.0, 0.45, n, 22, True, acc)
        add_text(s, 1.7, y + 0.18, 10.7, 0.4, title + ("   ←  my recommendation" if rec else ""), 18, True, NAVY)
        add_text(s, 1.7, y + 0.68, 10.7, 0.55, body, 14, False, SLATE)
        y += 1.55
    notes(s, "Stop talking. Let the advisor choose. If they want details, go to backup slides. If they choose 2, agree the next v2 experiment in writing before any new run.")

    # 19 Backup timeline
    s = blank(prs, 19, total, backup=True)
    header(s, "Backup", "Chronology (folder numbers are not time)", "Layer A → M0 freeze → K/U → papers → v2. Do not present folder names as the story.")
    tbl = s.shapes.add_table(8, 4, Inches(0.35), Inches(1.4), Inches(12.6), Inches(5.5)).table
    for i, w in enumerate([2.2, 3.3, 4.3, 2.8]):
        tbl.columns[i].width = Inches(w)
    data = [
        ["When", "Stage", "Result", "Status"],
        ["Aug 2026", "SVM dual-index MiniLM", "P@5 33% vs 36.5% word-count", "Superseded"],
        ["24–26 Aug", "Oracle pool → Method D", "68/78 = 87.18% Hit@5", "Frozen as M0"],
        ["27 Aug", "Phase 8 freeze + H labels", "H Success@5 25/40 diagnostic", "Burned"],
        ["Late Aug", "Phase 11 M0–M4", "All 68/78", "M0 kept"],
        ["Late Aug", "Phase 12 K / U / A2", "27/40 and 23/40; κ 0.55 / 0.68", "Official unseen M0"],
        ["6 Sep", "PLOS / IEEE packaging", "Submission snapshot", "No DOI in repo"],
        ["9–13 Sep", "ULTRA v2 Phases 2–8", "Dense 15/51; 3-way unsupported", "TRAIN/DEV only"],
    ]
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            set_cell(tbl, r, c, val, r == 0)
    style_table(tbl, True, 12)
    for r in range(1, 8):
        for c in range(4):
            for p in tbl.cell(r, c).text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if c != 0 else PP_ALIGN.CENTER
    notes(s, "Only if someone asks 'what was Phase 4?'. Historical Phase 4 chunk ANN is not v2 Phase 4 hybrid.")

    # 20 Backup talking points
    s = blank(prs, 20, total, backup=True)
    header(s, "Backup", "If challenged — exact sentences", "Use these; do not improvise a bigger number.")
    lines = [
        ("“Did you publish 87%?”", "“87.18% is the frozen development known-item Hit@5 on n=78. The unseen known-item figure is 67.5% on K. Usefulness is 57.5% on U.”"),
        ("“So the system failed?”", "“Urdu-script retrieval did not fail. Ordinary Roman Urdu failed. That is the result, and it is why v2 exists.”"),
        ("“Can you improve it before the thesis?”", "“Not on K/U — those sets are burned. v2 TEST is still sealed. Improving Method D after seeing K Roman 1/12 would not be an unseen test of M0.”"),
        ("“What is new vs other Urdu IR?”", "“A frozen script-aware BM25 protocol that separates development title_roman from sealed ordinary Roman, plus human Success@5 that is not called Hit@5.”"),
        ("“Why not neural as the official system?”", "“On the freeze pool, MiniLM lost to M0. On new Roman KN, dense is promising — that is v2, not the PLOS freeze.”"),
        ("“What do you want from me?”", "“Approve finishing the thesis on the three M0 numbers, and treat v2 as a follow-on with TEST unopened until we freeze one candidate.”"),
    ]
    y = 1.38
    for q, a in lines:
        add_text(s, 0.5, y, 12.3, 0.28, q, 13, True, TEAL)
        add_text(s, 0.5, y + 0.28, 12.3, 0.48, a, 13, False, SLATE)
        y += 0.88
    notes(s, "Keep this in reserve. The last Q is the close.")

    prs.save(OUT)
    print(f"Wrote {OUT} ({total} slides)")


if __name__ == "__main__":
    build()
